#!/usr/bin/env python3
"""Convert file.md to file.docx with strict formatting rules."""

import os
import re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = "Times New Roman"
FONT_SIZE = 12


def set_run_font(run, bold=False, italic=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE)
    run.bold = bold
    run.italic = italic
    run.font.highlight_color = None
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:shd", "w:rStyle"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)


def parse_inline(para, text):
    """Add runs to para parsing **bold**, *italic*, `code`."""
    pattern = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            set_run_font(para.add_run(text[pos:m.start()]))
        if m.group(2):
            set_run_font(para.add_run(m.group(2)), bold=True)
        elif m.group(3):
            set_run_font(para.add_run(m.group(3)), italic=True)
        elif m.group(4):
            set_run_font(para.add_run(m.group(4)), italic=True)
        pos = m.end()
    if pos < len(text):
        set_run_font(para.add_run(text[pos:]))


def fix_para_fonts(para):
    for run in para.runs:
        run.font.name = FONT_NAME
        if not run.font.size:
            run.font.size = Pt(FONT_SIZE)
        rPr = run._r.get_or_add_rPr()
        for tag in ("w:shd", "w:rStyle"):
            el = rPr.find(qn(tag))
            if el is not None:
                rPr.remove(el)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, bold=True)
    run.font.size = Pt(14 if level == 1 else 13 if level == 2 else 12)
    return p


def add_bullet_para(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    parse_inline(p, text)
    fix_para_fonts(p)
    return p


def add_numbered_para(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    parse_inline(p, text)
    fix_para_fonts(p)
    return p


def set_cell_shading_none(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        tcPr.remove(shd)


def get_list_number_abstract_num_id(doc):
    """Return the abstractNumId integer used by 'List Number' style."""
    style = doc.styles['List Number']
    numId_val = style.element.pPr.numPr.numId.get(qn('w:val'))
    numbering = doc.part.numbering_part._element
    for num in numbering.findall(qn('w:num')):
        if num.get(qn('w:numId')) == numId_val:
            return int(num.find(qn('w:abstractNumId')).get(qn('w:val')))
    return 1


def create_restarted_num_id(doc, abstract_num_id):
    """Add a new numId to numbering.xml that restarts at 1."""
    numbering = doc.part.numbering_part._element
    existing = numbering.findall(qn('w:num'))
    new_id = max(int(n.get(qn('w:numId'))) for n in existing) + 1

    num_el = OxmlElement('w:num')
    num_el.set(qn('w:numId'), str(new_id))

    abstract_ref = OxmlElement('w:abstractNumId')
    abstract_ref.set(qn('w:val'), str(abstract_num_id))
    num_el.append(abstract_ref)

    lvl_override = OxmlElement('w:lvlOverride')
    lvl_override.set(qn('w:ilvl'), '0')
    start_override = OxmlElement('w:startOverride')
    start_override.set(qn('w:val'), '1')
    lvl_override.append(start_override)
    num_el.append(lvl_override)

    numbering.append(num_el)
    return new_id


def apply_num_id_to_para(para, num_id):
    """Assign a specific numId to a paragraph's numPr."""
    pPr = para._p.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr')
        pPr.append(numPr)
    ilvl = numPr.find(qn('w:ilvl'))
    if ilvl is None:
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), '0')
        numPr.append(ilvl)
    numId_el = numPr.find(qn('w:numId'))
    if numId_el is None:
        numId_el = OxmlElement('w:numId')
        numPr.append(numId_el)
    numId_el.set(qn('w:val'), str(num_id))


# 1. text  — numbered item
LIST_NUM = re.compile(r'^(\d+)\.\s+(.*)')
# 1а. text — lettered sub-item (1а, 7б, 6в …)
LIST_LET = re.compile(r'^\d+[а-яёa-zA-Z]+\.\s+(.*)', re.IGNORECASE)


def parse_cell_content(doc, cell, raw_text, abstract_num_id):
    """Fill a table cell: 1./1а. items → numbered lists (per-cell restart)."""
    parts = re.split(r'<br\s*/?>', raw_text)

    for para in list(cell.paragraphs):
        para._p.getparent().remove(para._p)

    cell_num_id = None
    in_numbered = False

    for part in parts:
        part = part.strip()
        if not part:
            continue

        m_let = LIST_LET.match(part)
        m_num = LIST_NUM.match(part)

        if m_let or m_num:
            # Extract the actual text (group indices differ between the two patterns)
            item_text = (m_let.group(1) if m_let else m_num.group(2)).strip()

            if not in_numbered:
                cell_num_id = create_restarted_num_id(doc, abstract_num_id)
                in_numbered = True

            para = cell.add_paragraph(style=doc.styles['List Number'])
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(1)
            apply_num_id_to_para(para, cell_num_id)
            parse_inline(para, item_text)
        else:
            in_numbered = False
            para = cell.add_paragraph(style=doc.styles['Normal'])
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(1)
            parse_inline(para, part)

    for para in cell.paragraphs:
        fix_para_fonts(para)


def parse_markdown_table(doc, lines, start, abstract_num_id):
    rows = []
    i = start
    while i < len(lines):
        line = lines[i].strip()
        if not line.startswith('|'):
            break
        if re.match(r'^\|[\s\-:|]+\|', line) and re.search(r'-{2,}', line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows.append(cells)
        i += 1

    if not rows:
        return start + 1

    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'

    for ri, row_data in enumerate(rows):
        for ci, cell_text in enumerate(row_data):
            if ci >= ncols:
                break
            cell = table.rows[ri].cells[ci]
            set_cell_shading_none(cell)
            parse_cell_content(doc, cell, cell_text, abstract_num_id)

    for row in table.rows:
        for cell in row.cells:
            set_cell_shading_none(cell)

    doc.add_paragraph()
    return i


def convert(md_path, docx_path):
    doc = Document()

    for s_name in ['Normal', 'List Bullet', 'List Number']:
        try:
            s = doc.styles[s_name]
            s.font.name = FONT_NAME
            s.font.size = Pt(FONT_SIZE)
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(4)
        except Exception:
            pass

    abstract_num_id = get_list_number_abstract_num_id(doc)

    with open(md_path, encoding='utf-8') as f:
        lines = [line.rstrip('\n') for line in f]

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if re.match(r'^[-*_]{3,}$', stripped):
            i += 1
            continue

        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        if stripped.startswith('|'):
            i = parse_markdown_table(doc, lines, i, abstract_num_id)
            continue

        m = re.match(r'^[\*\-\+]\s+(.*)', line)
        if m:
            add_bullet_para(doc, m.group(1).strip())
            i += 1
            continue

        m = re.match(r'^\d+\.\s+(.*)', line)
        if m:
            add_numbered_para(doc, m.group(1).strip())
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        parse_inline(p, line)
        fix_para_fonts(p)
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    paths = [    ]
    for root, dirs, files in os.walk("task6/ru"):
        for file in files:
            if file.endswith(".md"):
                file_path = os.path.join(root, file)
                dest_path = os.path.join(root, file.replace(".md", ".docx"))
                print(f"Converting {file_path} to {dest_path}")
                convert(file_path, dest_path)
